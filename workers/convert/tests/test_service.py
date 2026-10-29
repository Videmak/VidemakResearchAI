"""The HTTP service around the Word package: LibreOffice conversions, PDF text, limits and safety."""
import base64
import io
import json
import os
import struct
import subprocess
import time
import zipfile

import pytest

import word_fixtures as fixtures
from conftest import needs_libreoffice, walk
from videmak_convert import config
from videmak_convert.jobs import DOCX_MIME, check_package
from videmak_convert.process import ConvertError, libreoffice_convert, run_bounded_process


@pytest.fixture(scope="module")
def compat():
    return fixtures.compat_docx()


def pdf_info(data: bytes):
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    try:
        texts = []
        for index in range(len(document)):
            page = document[index]
            textpage = page.get_textpage()
            texts.append(textpage.get_text_bounded())
            textpage.close()
            page.close()
        return len(texts), "\n".join(texts)
    finally:
        document.close()


def tiny_pdf(pages, title=None) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    count = len(pages)
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
               2: b"<< /Type /Pages /Kids [" + b" ".join(f"{4 + 2 * i} 0 R".encode() for i in range(count)) + f"] /Count {count} >>".encode(),
               3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    for index, text in enumerate(pages):
        page_id, content_id = 4 + 2 * index, 5 + 2 * index
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
        objects[page_id] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
                            f"/Contents {content_id} 0 R >>").encode()
        objects[content_id] = f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
    info_id = 4 + 2 * count
    if title:
        objects[info_id] = f"<< /Title ({title}) /Author (Ada Writer) >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = out.tell()
        out.write(f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n")
    xref = out.tell()
    size = max(objects) + 1
    out.write(f"xref\n0 {size}\n0000000000 65535 f \n".encode())
    for number in range(1, size):
        out.write(f"{offsets[number]:010d} 00000 n \n".encode() if number in offsets else b"0000000000 65535 f \n")
    trailer = f"<< /Size {size} /Root 1 0 R" + (f" /Info {info_id} 0 R" if title else "") + " >>"
    out.write(f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def tiny_wmf() -> bytes:
    """A placeable Windows metafile with one red rectangle."""
    records = [
        struct.pack("<IHhh", 5, 0x020B, 0, 0),  # SetWindowOrg(y, x)
        struct.pack("<IHhh", 5, 0x020C, 100, 200),  # SetWindowExt(y, x)
        struct.pack("<IHHBBBBH", 7, 0x02FC, 0, 220, 30, 30, 0, 0),  # CreateBrushIndirect(solid red)
        struct.pack("<IHH", 4, 0x012D, 0),  # SelectObject(0)
        struct.pack("<IHhhhh", 7, 0x041B, 90, 190, 10, 10),  # Rectangle(bottom, right, top, left)
        struct.pack("<IH", 3, 0x0000),  # EOF
    ]
    body = b"".join(records)
    header = struct.pack("<HHHIHIH", 1, 9, 0x0300, 9 + len(body) // 2, 1, 7, 0)
    placeable = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, 200, 100, 1440, 0)
    checksum = 0
    for (word,) in struct.iter_unpack("<H", placeable):
        checksum ^= word
    return placeable + struct.pack("<H", checksum) + header + body


def docx_with_wmf() -> bytes:
    """compat-style document whose picture is a WMF instead of a PNG."""
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("A chart drawn in an old Windows format:")
    document.add_paragraph().add_run().add_picture(io.BytesIO(fixtures.png()), width=Inches(2))
    buffer = io.BytesIO()
    document.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            name = item.filename
            if name.startswith("word/media/") and name.endswith(".png"):
                name, data = name[:-4] + ".wmf", tiny_wmf()
            elif name == "word/_rels/document.xml.rels":
                data = data.replace(b'.png"', b'.wmf"')
            elif name == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Default Extension="wmf" ContentType="image/x-wmf"/></Types>')
            target.writestr(name, data)
    return output.getvalue()


# ------------------------------------------------------------------ health


def test_health(client):
    response = client.get("/health")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok" and body["service"] == "videmak-convert" and body["version"]
    assert set(body) == {"status", "service", "version", "libreoffice", "workers", "busy"}


# ------------------------------------------------------------------ LibreOffice formats


@needs_libreoffice
@pytest.mark.parametrize("extension", ["doc", "odt"])
def test_import_converts_doc_and_odt_through_libreoffice(client, compat, extension):
    data = libreoffice_convert(compat, "docx", extension, timeout=120)
    response = client.post("/word/import", files={"file": (f"thesis.{extension}", data, "application/octet-stream")})
    assert response.status_code == 200, response.text
    body = response.json()
    # LibreOffice joins the fixture's two back-to-back tables into one when it saves .doc/.odt.
    assert body["title"] == "thesis" and body["stats"]["tables"] >= 1 and body["stats"]["headings"] >= 4
    headings = [b for b in walk(body["blocks"]) if b["type"] == "heading"]
    assert any("Methods" in json.dumps(b["content"]) for b in headings)
    assert any(b["type"] == "image" for b in walk(body["blocks"])) and len(body["assets"]) == 1
    source = body["sourceDocx"]
    assert source["format"] == extension and source["converted"] is True and source["contentType"] == DOCX_MIME
    template = base64.b64decode(source["data"])
    assert len(template) == source["size"]
    with zipfile.ZipFile(io.BytesIO(template)) as package:
        assert "word/document.xml" in package.namelist()


@needs_libreoffice
def test_import_converts_rtf(client):
    rtf = (rb"{\rtf1\ansi\deff0{\fonttbl{\f0\froman Times New Roman;}}"
           rb"{\pard\b Heading in RTF\b0\par}{\pard Plain text with {\i italics}.\par}}")
    response = client.post("/word/import", files={"file": ("notes.rtf", rtf, "application/rtf")})
    assert response.status_code == 200, response.text
    body = response.json()
    text = json.dumps(body["blocks"])
    assert "Heading in RTF" in text and "italics" in text and body["title"] == "notes"
    assert body["sourceDocx"]["converted"] is True and body["sourceDocx"]["data"]


@needs_libreoffice
def test_windows_metafile_images_become_png(client):
    response = client.post("/word/import", files={"file": ("chart.docx", docx_with_wmf(), DOCX_MIME)})
    assert response.status_code == 200, response.text
    body = response.json()
    (asset,) = body["assets"]
    assert asset["contentType"] == "image/png" and asset["name"].endswith(".png")
    assert base64.b64decode(asset["data"]).startswith(b"\x89PNG"), body["warnings"]


@needs_libreoffice
def test_export_as_pdf(client):
    blocks = [{"id": "h", "type": "heading", "props": {"level": 1}, "content": [{"type": "text", "text": "Introduction", "styles": {}}], "children": []},
              {"id": "p", "type": "paragraph", "props": {}, "content": [{"type": "text", "text": "Malaria in pregnancy.", "styles": {}}], "children": []},
              {"id": "b", "type": "pageBreak", "props": {}, "children": []},
              {"id": "q", "type": "paragraph", "props": {}, "content": [{"type": "text", "text": "Second page.", "styles": {}}], "children": []}]
    response = client.post("/word/export", json={"title": "Report é", "blocks": blocks, "format": "pdf"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf" and response.content.startswith(b"%PDF")
    assert "filename*=UTF-8''Report%20%C3%A9.pdf" in response.headers["content-disposition"]
    pages, text = pdf_info(response.content)
    assert pages == 2 and "Introduction" in text and "Second page." in text


@needs_libreoffice
def test_docx_to_pdf(client, compat):
    response = client.post("/convert/docx-to-pdf", files={"file": ("thesis.docx", compat, DOCX_MIME)})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert 'filename="thesis.pdf"' in response.headers["content-disposition"]
    pages, text = pdf_info(response.content)
    assert pages >= 2 and "Treatment Adherence" in text and "Running head" in text
    assert client.post("/convert/docx-to-pdf", files={"file": ("notes.txt", b"hello", "text/plain")}).status_code == 400
    empty = client.post("/convert/docx-to-pdf", files={"file": ("empty.docx", b"", DOCX_MIME)})
    assert empty.status_code == 400 and empty.json()["detail"] == "The editor document is empty"
    # LibreOffice would print a non-Word file as plain text; a .docx must be a Word package.
    broken = client.post("/convert/docx-to-pdf", files={"file": ("broken.docx", b"not a zip at all", DOCX_MIME)})
    assert broken.status_code == 400 and broken.json()["detail"] == "This file could not be read as a Word document."


# ------------------------------------------------------------------ PDF text


def test_pdf_text_per_page(client):
    data = tiny_pdf(["Hello first page", "And the second page"], title="A Tiny Paper")
    response = client.post("/pdf/text", files={"file": ("paper.pdf", data, "application/pdf")})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["page_count"] == 2 and [page["page"] for page in body["pages"]] == [1, 2]
    assert body["pages"][0]["text"].strip() == "Hello first page"
    assert body["pages"][1]["text"].strip() == "And the second page"
    assert body["title"] == "A Tiny Paper" and body["metadata"]["Author"] == "Ada Writer"
    untitled = client.post("/pdf/text", files={"file": ("x.pdf", tiny_pdf(["Only"]), "application/pdf")}).json()
    assert untitled["title"] is None and untitled["page_count"] == 1 and isinstance(untitled["metadata"], dict)


@needs_libreoffice
def test_pdf_text_of_a_real_document(client, compat):
    pdf = libreoffice_convert(compat, "docx", "pdf", timeout=120)
    body = client.post("/pdf/text", files={"file": ("thesis.pdf", pdf, "application/pdf")}).json()
    assert body["page_count"] >= 2 and len(body["pages"]) == body["page_count"]
    assert "Treatment Adherence" in body["pages"][0]["text"] and "\r" not in body["pages"][0]["text"]
    assert "Results" in "".join(page["text"] for page in body["pages"][1:])


@needs_libreoffice
def test_pdf_text_refuses_password_protected_pdfs(client, compat):
    options = json.dumps({"EncryptFile": {"type": "boolean", "value": "true"},
                          "DocumentOpenPassword": {"type": "string", "value": "secret"}})
    import videmak_convert.process as process

    original = process.FILTERS["pdf"]
    process.FILTERS["pdf"] = "pdf:writer_pdf_Export:" + options
    try:
        locked = libreoffice_convert(compat, "docx", "pdf", timeout=120)
    finally:
        process.FILTERS["pdf"] = original
    response = client.post("/pdf/text", files={"file": ("locked.pdf", locked, "application/pdf")})
    assert response.status_code == 400 and "password-protected" in response.json()["detail"]


def test_pdf_text_rejects_other_files(client):
    assert client.post("/pdf/text", files={"file": ("a.pdf", b"", "application/pdf")}).status_code == 400
    not_pdf = client.post("/pdf/text", files={"file": ("a.pdf", b"PK\x03\x04 not a pdf", "application/pdf")})
    assert not_pdf.status_code == 400 and not_pdf.json()["detail"] == "Upload a PDF file."
    broken = client.post("/pdf/text", files={"file": ("a.pdf", b"%PDF-1.7\n garbage", "application/pdf")})
    assert broken.status_code == 400 and broken.json()["detail"] == "This file could not be read as a PDF."


# ------------------------------------------------------------------ limits and safety


def test_size_limits(client, monkeypatch, compat):
    monkeypatch.setattr(config, "MAX_IMPORT_BYTES", 1000)
    monkeypatch.setattr(config, "MAX_EXPORT_BYTES", 1000)
    big = b"x" * 2000
    response = client.post("/word/import", files={"file": ("big.docx", big, DOCX_MIME)})
    assert response.status_code == 413 and response.json()["detail"] == "Files up to 50 MB can be opened"
    response = client.post("/convert/docx-to-pdf", files={"file": ("big.docx", big, DOCX_MIME)})
    assert response.status_code == 413 and response.json()["detail"] == "Editor exports must be 50 MB or smaller"
    assert client.post("/pdf/text", files={"file": ("big.pdf", b"%PDF-" + big, "application/pdf")}).status_code == 413
    response = client.post("/word/export", json={"title": "T", "blocks": [{"type": "paragraph", "content": [{"type": "text", "text": "y" * 2000}]}]})
    assert response.status_code == 413
    response = client.post("/word/export", data={"payload": json.dumps({"title": "T"})}, files={"source": ("s.docx", compat, DOCX_MIME)})
    assert response.status_code == 413 and "source" in response.json()["detail"].lower()

    # Whole requests over the path's limit are refused before they are read...
    monkeypatch.setattr(config, "MAX_IMPORT_BYTES", 0)
    too_big = b"z" * (2 * 1024 * 1024)
    response = client.post("/word/import", files={"file": ("big.docx", too_big, DOCX_MIME)})
    assert response.status_code == 413 and response.json()["detail"] == "The request is too large."

    # ...and so are streamed (chunked) ones without a Content-Length.
    def chunks():
        for _ in range(4):
            yield b"z" * (1024 * 1024)

    response = client.post("/word/import", content=chunks(), headers={"Content-Type": "multipart/form-data; boundary=xyz"})
    assert response.status_code == 413


def test_unpacked_size_guard(compat):
    check_package(compat)
    with pytest.raises(ConvertError) as error:
        check_package(compat, max_unpacked=1000)
    assert error.value.status == 413
    with pytest.raises(ConvertError) as error:
        check_package(b"PK\x03\x04 broken")
    assert error.value.status == 400
    no_document = io.BytesIO()
    with zipfile.ZipFile(no_document, "w") as package:
        package.writestr("hello.txt", "hi")
    with pytest.raises(ConvertError):
        check_package(no_document.getvalue())


def test_export_validation(client, monkeypatch):
    assert client.post("/word/export", json={"format": "html"}).status_code == 422
    assert client.post("/word/export", content=b"{not json", headers={"Content-Type": "application/json"}).status_code == 400
    assert client.post("/word/export", json=[1, 2]).status_code == 400
    assert client.post("/word/export", json={"source_base64": "***"}).status_code == 400
    assert client.post("/word/export", json={"asset_fetch_base": "ftp://x/"}).status_code == 400
    assert client.post("/word/export", data={"other": "x"}, files={"source": ("s.docx", b"x", DOCX_MIME)}).status_code == 400
    monkeypatch.setattr(config, "ASSET_FETCH_BASES", ["http://videmak-api:4000/internal/assets/"])
    refused = client.post("/word/export", json={"asset_fetch_base": "http://elsewhere:80/"})
    assert refused.status_code == 400 and refused.json()["detail"] == "asset_fetch_base is not allowed here"
    allowed = client.post("/word/export", json={"asset_fetch_base": "http://videmak-api:4000/internal/assets/", "blocks": []})
    assert allowed.status_code == 200


def test_busy_converter_answers_503(client, monkeypatch, compat):
    monkeypatch.setattr(config, "MAX_QUEUE", -1000)
    response = client.post("/word/import", files={"file": ("thesis.docx", compat, DOCX_MIME)})
    assert response.status_code == 503 and response.headers["retry-after"] == "5"


def test_bounded_process_ends_the_whole_process_group():
    # A child (sleep 3071) and a grandchild (sleep 3072): both must be gone after the timeout.
    started = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        run_bounded_process(["sh", "-c", "sleep 3072 & sleep 3071"], timeout=1)
    assert time.monotonic() - started < 10

    def alive():
        found = []
        for pid in filter(str.isdigit, os.listdir("/proc")):
            try:
                with open(f"/proc/{pid}/cmdline", "rb") as handle:
                    command = handle.read().replace(b"\0", b" ").decode(errors="replace").strip()
            except OSError:
                continue
            if command in ("sleep 3071", "sleep 3072"):
                found.append(command)
        return found

    time.sleep(0.5)
    assert alive() == []
