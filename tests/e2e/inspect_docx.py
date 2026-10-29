"""What a Word file contains, as Word would see it: text with styles, runs' formatting, lists,
tables, images, footnotes, comments, tracked changes, headers/footers, page setup."""
import sys
import zipfile

from docx import Document
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
path = sys.argv[1]
doc = Document(path)
section = doc.sections[0]
print(f"page {section.page_width.inches:.2f}x{section.page_height.inches:.2f} in, margins L{section.left_margin.inches:.2f} R{section.right_margin.inches:.2f} T{section.top_margin.inches:.2f}")
print("header:", repr(" | ".join(p.text for p in section.header.paragraphs if p.text)), "| footer:", repr(" | ".join(p.text for p in section.footer.paragraphs if p.text)))
with zipfile.ZipFile(path) as package:
    names = package.namelist()
    print("parts:", ", ".join(n for n in names if n.startswith("word/") and not n.startswith("word/theme")))
    body = etree.fromstring(package.read("word/document.xml"))
    footer_fields = []
    for name in names:
        if name.startswith("word/footer") or name.startswith("word/header"):
            xml = package.read(name).decode()
            footer_fields += [f for f in ("PAGE", "NUMPAGES") if f" {f} " in xml or f">{f}<" in xml or f"{f} \\" in xml]
    print("page number fields in headers/footers:", footer_fields)
    print("footnotes:", len(etree.fromstring(package.read("word/footnotes.xml")).findall(W + "footnote")) - 2 if "word/footnotes.xml" in names else 0,
          "| comments:", len(etree.fromstring(package.read("word/comments.xml")).findall(W + "comment")) if "word/comments.xml" in names else 0,
          "| tracked ins/del:", len(body.findall(".//" + W + "ins")), len(body.findall(".//" + W + "del")),
          "| math:", len(body.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")),
          "| images:", len(body.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")),
          "| fields:", sorted({(i.text or '').strip().split(' ')[0] for i in body.iter(W + "instrText") if (i.text or '').strip()}))
for paragraph in doc.paragraphs:
    if not paragraph.text.strip():
        continue
    runs = []
    for run in paragraph.runs:
        flags = [n for n, v in (("b", run.bold), ("i", run.italic), ("u", run.underline), ("s", run.font.strike), ("sup", run.font.superscript), ("sub", run.font.subscript)) if v]
        if run.font.color is not None and run.font.color.rgb is not None and str(run.font.color.rgb) not in ("000000",):
            flags.append(f"#{run.font.color.rgb}")
        if run.font.highlight_color:
            flags.append(f"hl:{run.font.highlight_color}")
        if run.font.name:
            flags.append(run.font.name)
        if run.font.size:
            flags.append(f"{run.font.size.pt:g}pt")
        if flags and run.text.strip():
            runs.append(f"{run.text.strip()[:18]}[{','.join(flags)}]")
    fmt = paragraph.paragraph_format
    extra = []
    if fmt.left_indent:
        extra.append(f"indent {fmt.left_indent.pt:g}pt")
    if fmt.first_line_indent:
        extra.append(f"first {fmt.first_line_indent.pt:g}pt")
    if fmt.line_spacing:
        extra.append(f"line {fmt.line_spacing}")
    if fmt.space_before:
        extra.append(f"before {fmt.space_before.pt:g}pt")
    if paragraph.alignment is not None:
        extra.append(f"align {paragraph.alignment}")
    print(f"- [{paragraph.style.name}] {paragraph.text[:70]!r} {' '.join(extra)} {' '.join(runs[:8])}")
for index, table in enumerate(doc.tables):
    print(f"table {index + 1}: {len(table.rows)}x{len(table.columns)}:", [[cell.text for cell in row.cells] for row in table.rows])
